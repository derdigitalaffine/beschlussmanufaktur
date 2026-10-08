"""Local rendering only. Text, tables and lists; no URL, image or external asset retrieval."""
import io,re
from html import escape
from html.parser import HTMLParser
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib.pagesizes import A4
from docx import Document
from .templates_service import render_markdown

class Blocks(HTMLParser):
    def __init__(self):super().__init__();self.blocks=[];self.text='';self.tag='p';self.rows=None;self.row=None;self.cell=None
    def flush(self):
        if self.text.strip():self.blocks.append((self.tag,self.text.strip()))
        self.text=''
    def handle_starttag(self,tag,attrs):
        if tag=='table':self.flush();self.rows=[]
        elif tag=='tr' and self.rows is not None:self.row=[]
        elif tag in ('td','th') and self.rows is not None:self.cell=''
        elif self.rows is None and tag in ('p','h1','h2','h3','h4','li','blockquote','pre'):self.flush();self.tag=tag
        elif tag=='br':self.handle_data('\n')
    def handle_endtag(self,tag):
        if tag in ('td','th') and self.cell is not None:self.row.append(self.cell.strip());self.cell=None
        elif tag=='tr' and self.row is not None:self.rows.append(self.row);self.row=None
        elif tag=='table':self.blocks.append(('table',self.rows));self.rows=None
        elif self.rows is None and tag in ('p','h1','h2','h3','h4','li','blockquote','pre'):self.flush();self.tag='p'
    def handle_data(self,data):
        if self.cell is not None:self.cell+=data
        elif self.rows is None:self.text+=data

def blocks(markdown):
    parser=Blocks();parser.feed(render_markdown(markdown));parser.flush();return parser.blocks

def export(title,markdown,format='pdf',subtitle='',metadata=None,organization_id=None):
    from .models import PortalConfiguration
    brand=PortalConfiguration.objects.filter(organization_id=organization_id).first() if organization_id else None
    brand_name=brand.title if brand else 'Beschlussmanufaktur'
    brand_color=brand.color if brand and re.fullmatch(r'#[0-9a-fA-F]{6}',brand.color) else '#245846'
    parts=blocks(markdown)
    if format=='docx':
        doc=Document()
        from docx.shared import RGBColor
        doc.sections[0].header.paragraphs[0].text=brand_name
        doc.sections[0].footer.paragraphs[0].text=brand_name
        for name in ('Title','Heading 1','Heading 2','Heading 3','Heading 4'):doc.styles[name].font.color.rgb=RGBColor.from_string(brand_color[1:])
        doc.add_heading(title,0)
        if subtitle:doc.add_paragraph(subtitle)
        for tag,text in parts:
            if tag=='table':
                if not text:continue
                table=doc.add_table(rows=len(text),cols=max(len(row) for row in text));table.style='Table Grid'
                for i,row in enumerate(text):
                    for j,cell in enumerate(row):table.cell(i,j).text=cell
            elif tag.startswith('h'):doc.add_heading(text,min(int(tag[1]),4))
            else:doc.add_paragraph(text,style='List Bullet' if tag=='li' else None)
        if metadata:
            doc.add_heading('Anlagen / Stand',1)
            for line in metadata:doc.add_paragraph(line)
        buffer=io.BytesIO();doc.save(buffer);return buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format!='pdf':raise ValueError('Unbekanntes Exportformat')
    styles=getSampleStyleSheet();styles['Normal'].leading=15
    for name in ('Title','Heading1','Heading2','Heading3'):styles[name].textColor=colors.HexColor(brand_color)
    elements=[Paragraph(escape(title),styles['Title'])]
    if subtitle:elements.append(Paragraph(escape(subtitle),styles['Normal']))
    elements.append(Spacer(1,14))
    for tag,text in parts:
        if tag=='table':
            if not text:continue
            cols=max(len(row) for row in text)
            rows=[[Paragraph(escape(cell),styles['Normal']) for cell in row]+['']*(cols-len(row)) for row in text]
            table=Table(rows,colWidths=[(A4[0]-88)/cols]*cols,repeatRows=1)
            table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf3f0')),('VALIGN',(0,0),(-1,-1),'TOP')]))
            elements.append(table)
        else:
            style=styles['Heading'+str(min(int(tag[1]),3))] if tag.startswith('h') else styles['Normal']
            content=escape(('• ' if tag=='li' else '')+text).replace('\n','<br/>')
            elements.extend([Paragraph(content,style),Spacer(1,8)])
    if metadata:
        elements.append(Paragraph('Anlagen / Stand',styles['Heading2']))
        for line in metadata:elements.append(Paragraph(escape(line),styles['Normal']))
    buffer=io.BytesIO()
    def footer(canvas,doc):canvas.saveState();canvas.setFont('Helvetica',9);canvas.drawString(44,24,brand_name);canvas.drawRightString(A4[0]-44,24,f'Seite {doc.page}');canvas.restoreState()
    SimpleDocTemplate(buffer,pagesize=A4,leftMargin=44,rightMargin=44,topMargin=44,bottomMargin=44,title=title,author='Beschlussmanufaktur').build(elements,onFirstPage=footer,onLaterPages=footer)
    return buffer.getvalue(),'application/pdf'
