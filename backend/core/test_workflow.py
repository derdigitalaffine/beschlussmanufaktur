import io
from django.test import TestCase
from django.core.exceptions import ValidationError,PermissionDenied
from .models import Organization,User,Membership,TemplateKind,Template,ReviewStep,Publication,TemplateVersion,Attachment,Consultation,RegistryRecord
from .templates_service import archive,touch,template_access
from .workflow import submit,decide,publish
from .document_export import export
from docx import Document

class WorkflowTests(TestCase):
    def setUp(self):
        self.org=Organization.objects.create(name='VG',kind='association')
        self.author=User.objects.create_user(username='author',email='author@example.org')
        self.reviewer=User.objects.create_user(username='reviewer',email='review@example.org')
        self.clerk=User.objects.create_user(username='clerk',email='clerk@example.org')
        self.a=Membership.objects.create(user=self.author,organization=self.org,role='author')
        self.r=Membership.objects.create(user=self.reviewer,organization=self.org,role='reviewer')
        self.c=Membership.objects.create(user=self.clerk,organization=self.org,role='clerk')
        self.kind=TemplateKind.objects.create(organization=self.org,name='Beschluss',four_eyes=True)
        self.obj=Template.objects.create(organization=self.org,kind=self.kind,author=self.author,subject='Bau',markdown='Nicht öffentlich',public_markdown='Öffentlich geprüft',classification='public_planned')
        archive(self.obj,self.author,'Entwurf')
    def ready(self):
        submit(self.a,self.obj.pk,1)
        steps=list(ReviewStep.objects.order_by('group'))
        decide(self.r,steps[0].pk,1,'approve','Fachlich richtig');decide(self.c,steps[1].pk,1,'approve','Bereitgestellt')
        self.obj.refresh_from_db()
    def test_linear_default_number_at_readiness(self):
        self.assertFalse(self.obj.number);self.ready();self.assertEqual(self.obj.state,'ready');self.assertTrue(self.obj.number.endswith('/0001'))
    def test_later_step_cannot_skip_first(self):
        submit(self.a,self.obj.pk,1);step=ReviewStep.objects.get(role='clerk')
        with self.assertRaises(ValidationError):decide(self.c,step.pk,1,'approve','')
    def test_reviewer_gets_only_current_assigned_version(self):
        self.assertFalse(template_access(self.r,'read',self.obj));submit(self.a,self.obj.pk,1)
        self.obj.refresh_from_db();self.assertTrue(template_access(self.r,'read',self.obj));self.assertFalse(template_access(self.r,'edit',self.obj))
    def test_return_requires_comment(self):
        submit(self.a,self.obj.pk,1);step=ReviewStep.objects.get(role='reviewer')
        with self.assertRaises(ValidationError):decide(self.r,step.pk,1,'return','')
        decide(self.r,step.pk,1,'return','Bitte überarbeiten');self.obj.refresh_from_db();self.assertEqual(self.obj.state,'returned')
    def test_changes_invalidate_review(self):
        submit(self.a,self.obj.pk,1);step=ReviewStep.objects.get(role='reviewer')
        self.obj.refresh_from_db();self.obj.markdown='Neu';touch(self.obj,self.author,'Fachliche Änderung')
        with self.assertRaises(PermissionDenied):decide(self.r,step.pk,1,'approve','')
    def test_parallel_conditional_groups(self):
        self.kind.fields=[{'key':'finance','label':'Finanzen','type':'boolean'}]
        self.kind.workflow=[{'name':'Fachlich','role':'reviewer','group':1},{'name':'Finanzen','role':'release','group':1,'condition':{'field':'finance','equals':True}},{'name':'Bereitstellen','role':'clerk','group':2}];self.kind.save()
        self.obj.fields={'finance':False};self.obj.save();submit(self.a,self.obj.pk,1)
        self.assertEqual(ReviewStep.objects.count(),2)
    def test_four_eyes_prevents_author_approval(self):
        context=Membership.objects.create(user=self.author,organization=self.org,role='reviewer')
        submit(self.a,self.obj.pk,1)
        with self.assertRaises(ValidationError):decide(context,ReviewStep.objects.get(role='reviewer').pk,1,'approve','')
    def test_publication_requires_ready_public_text(self):
        with self.assertRaises(ValidationError):publish(self.c,self.obj.pk,1)
        self.ready();publish(self.c,self.obj.pk,1);publication=Publication.objects.get();self.assertEqual(publication.markdown,'Öffentlich geprüft');self.assertNotIn('Nicht öffentlich',publication.markdown)
        publish(self.c,self.obj.pk,1,withdraw=True);publication.refresh_from_db();self.assertIsNotNone(publication.withdrawn_at)
    def test_number_retained_after_revision(self):
        self.ready();number=self.obj.number;self.obj.markdown='Änderung';touch(self.obj,self.author,'Änderung');self.assertEqual(self.obj.number,number);self.assertEqual(self.obj.state,'draft')
    def test_export_pdf_docx_no_network_images(self):
        md='# Überschrift\n\nText mit Umlauten äöü.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n![Bild](http://127.0.0.1/secret)'
        pdf,mime=export('Vorlage',md,'pdf');self.assertTrue(pdf.startswith(b'%PDF'));self.assertEqual(mime,'application/pdf')
        docx,mime=export('Vorlage',md,'docx');document=Document(io.BytesIO(docx));self.assertEqual(document.tables[0].cell(1,1).text,'2')
    def test_unchecked_attachment_blocks_ready(self):
        Attachment.objects.create(template=self.obj,name='Test.txt',file='private.txt',digest='a'*64,size=1,media_type='text/plain')
        submit(self.a,self.obj.pk,1);steps=list(ReviewStep.objects.order_by('group'));decide(self.r,steps[0].pk,1,'approve','')
        with self.assertRaises(ValidationError):decide(self.c,steps[1].pk,1,'approve','')
        steps[1].refresh_from_db();self.assertEqual(steps[1].state,'pending')
