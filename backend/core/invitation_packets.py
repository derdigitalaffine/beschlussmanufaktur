import io,hashlib
from datetime import UTC
from django.core.exceptions import PermissionDenied,ValidationError
from pypdf import PdfReader,PdfWriter
from reportlab.pdfgen import canvas
from PIL import Image
from docx import Document
from .models import Attachment,Template
from .templates_service import template_access
from .document_export import export
from .meetings_service import paper_markdown

def calendar(data):
    from django.utils.dateparse import parse_datetime
    def esc(value):return str(value).replace('\\','\\\\').replace('\r','').replace('\n','\\n').replace(';','\\;').replace(',','\\,')
    def utc(value):return parse_datetime(value).astimezone(UTC).strftime('%Y%m%dT%H%M%SZ')
    lines=['BEGIN:VCALENDAR','VERSION:2.0','PRODID:-//Beschlussmanufaktur//Sitzung//DE','BEGIN:VEVENT','UID:'+esc(data['id'])+'@beschlussmanufaktur','DTSTAMP:'+utc(data['starts_at']),'DTSTART:'+utc(data['starts_at']),'DTEND:'+utc(data['ends_at']),'SUMMARY:'+esc(data['title']),'LOCATION:'+esc(data['location']),'END:VEVENT','END:VCALENDAR']
    # Fold by UTF-8 byte lengths; CRLF injection is escaped above.
    folded=[]
    for line in lines:
        current=''
        for char in line:
            if len((current+char).encode())>73:folded.append(current);current=' '+char
            else:current+=char
        folded.append(current)
    return ('\r\n'.join(folded)+'\r\n').encode()

def attachment_pdf(binary,name,media):
    if media=='application/pdf':
        reader=PdfReader(io.BytesIO(binary))
        if reader.is_encrypted:raise ValidationError('Verschlüsselte PDF kann nicht in die Gesamtmappe aufgenommen werden.')
        if len(reader.pages)>500:raise ValidationError('Anlage hat mehr als 500 Seiten.')
        return binary
    if media=='text/plain':text=binary.decode('utf-8')
    elif 'wordprocessingml' in media:
        doc=Document(io.BytesIO(binary));text='\n\n'.join(p.text for p in doc.paragraphs)
        for table in doc.tables:text+='\n\n'+'\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
    elif media in ('image/png','image/jpeg'):
        image=Image.open(io.BytesIO(binary))
        if image.width*image.height>15000000:raise ValidationError('Bild überschreitet die erlaubte Pixelzahl.')
        output=io.BytesIO();image.convert('RGB').save(output,format='PDF');return output.getvalue()
    else:raise ValidationError('Anlage kann nicht als PDF ausgegeben werden.')
    return export(name,text,'pdf')[0]

def packet(data,context=None,asset_loader=None):
    files=[];seen=set();metadata=[]
    for item in data['items']:
        template=item.get('template')
        if not template:continue
        if context:
            obj=Template.objects.get(pk=template['id'])
            if not template_access(context,'export',obj):raise PermissionDenied('Aktuelle Rechte erlauben diese archivierte Unterlage nicht mehr.')
        for meta in template['attachments']:
            if meta['id'] in seen:continue
            seen.add(meta['id'])
            if not meta['checked']:raise ValidationError('Ungeprüfte Anlage im Versandstand.')
            if asset_loader:binary=asset_loader(meta)
            else:
                obj=Attachment.objects.get(pk=meta['id'],digest=meta['digest'])
                with obj.file.open('rb') as file:binary=file.read()
            if hashlib.sha256(binary).hexdigest()!=meta['digest']:raise ValidationError('Prüfsumme einer Anlage stimmt nicht.')
            converted=attachment_pdf(binary,meta['name'],meta['media_type'])
            files.append((meta['name'],converted,len(PdfReader(io.BytesIO(converted)).pages)))
    text=paper_markdown(data)
    base=export(data['title'],text,'pdf')[0]
    # Resolve page references after rendering; iterate until the cover page count stabilizes.
    for attempt in range(4):
        base_count=len(PdfReader(io.BytesIO(base)).pages);start=base_count+1
        metadata=[]
        for name,binary,count in files:metadata.append(f'{name} · Seiten {start}–{start+count-1}');start+=count
        updated=export(data['title'],text,'pdf',metadata=metadata)[0]
        if len(PdfReader(io.BytesIO(updated)).pages)==base_count:base=updated;break
        base=updated
    writer=PdfWriter()
    for binary in [base]+[f[1] for f in files]:
        reader=PdfReader(io.BytesIO(binary));writer.append(reader,import_outline=False)
    total=len(writer.pages)
    if total>2000:raise ValidationError('Gesamtmappe überschreitet 2000 Seiten.')
    for index,page in enumerate(writer.pages):
        # Remove document actions from assembled copies; original files remain unchanged.
        for action in ['/AA','/A']:page.pop(action,None)
        page.pop('/Annots',None)
        stamp=io.BytesIO();c=canvas.Canvas(stamp,pagesize=(float(page.mediabox.width),float(page.mediabox.height)))
        c.setFillColorRGB(1,1,1);c.rect(0,10,float(page.mediabox.width),24,fill=1,stroke=0);c.setFillColorRGB(0,0,0);c.setFont('Helvetica',9);c.drawRightString(float(page.mediabox.width)-35,20,f'Seite {index+1} / {total}');c.save()
        page.merge_page(PdfReader(io.BytesIO(stamp.getvalue())).pages[0])
    for key in ['/OpenAction','/AA','/AcroForm','/Names']:writer._root_object.pop(key,None)
    writer.add_metadata({'/Title':data['title'],'/Author':'Beschlussmanufaktur'})
    output=io.BytesIO();writer.write(output);return output.getvalue()
